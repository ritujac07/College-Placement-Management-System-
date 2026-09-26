import tkinter as tk
from tkinter import messagebox
import openpyxl
import os

FILE = "placement_data.xlsx"

# Create Excel file if it does not exist
if not os.path.exists(FILE):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([
        "Student ID",
        "Student Name",
        "Branch",
        "Year",
        "CGPA",
        "Company Name",
        "Placement Status"
    ])
    wb.save(FILE)


# ---------------- ADD STUDENT ----------------
def add_student():
    form = tk.Toplevel(dashboard)
    form.title("Add Student")
    form.geometry("400x550")

    tk.Label(
        form,
        text="Add Student",
        font=("Arial", 18, "bold")
    ).pack(pady=15)

    labels = [
        "Student ID",
        "Student Name",
        "Branch",
        "Year",
        "CGPA",
        "Company Name",
        "Placement Status"
    ]

    entries = {}

    for label in labels:
        tk.Label(form, text=label).pack()
        entry = tk.Entry(form)
        entry.pack(pady=4)
        entries[label] = entry

    def save_student():
        values = []

        for label in labels:
            value = entries[label].get().strip()

            if value == "":
                messagebox.showerror(
                    "Error",
                    "Please fill all fields."
                )
                return

            values.append(value)

        wb = openpyxl.load_workbook(FILE)
        ws = wb.active

        # Check duplicate Student ID
        for row in ws.iter_rows(min_row=2, values_only=True):
            if str(row[0]) == values[0]:
                messagebox.showerror(
                    "Error",
                    "Student ID already exists."
                )
                wb.close()
                return

        ws.append(values)
        wb.save(FILE)
        wb.close()

        messagebox.showinfo(
            "Success",
            "Student added successfully!"
        )

        form.destroy()

    tk.Button(
        form,
        text="Save Student",
        width=20,
        command=save_student
    ).pack(pady=20)


# ---------------- VIEW RECORDS ----------------
def view_records():
    window = tk.Toplevel(dashboard)
    window.title("View Records")
    window.geometry("800x500")

    tk.Label(
        window,
        text="Student Records",
        font=("Arial", 18, "bold")
    ).pack(pady=10)

    text = tk.Text(window, width=95, height=25)
    text.pack(padx=10, pady=10)

    wb = openpyxl.load_workbook(FILE)
    ws = wb.active

    for row in ws.iter_rows(values_only=True):
        text.insert(tk.END, " | ".join(str(x) for x in row) + "\n")

    wb.close()


# ---------------- SEARCH STUDENT ----------------
def search_student():
    window = tk.Toplevel(dashboard)
    window.title("Search Student")
    window.geometry("450x350")

    tk.Label(
        window,
        text="Search by Student ID",
        font=("Arial", 16, "bold")
    ).pack(pady=20)

    entry = tk.Entry(window)
    entry.pack(pady=10)

    result = tk.Text(window, width=50, height=10)
    result.pack(pady=10)

    def search():
        student_id = entry.get().strip()

        wb = openpyxl.load_workbook(FILE)
        ws = wb.active

        found = False

        for row in ws.iter_rows(min_row=2, values_only=True):
            if str(row[0]) == student_id:
                result.delete("1.0", tk.END)
                result.insert(
                    tk.END,
                    "Student ID: " + str(row[0]) + "\n"
                    "Name: " + str(row[1]) + "\n"
                    "Branch: " + str(row[2]) + "\n"
                    "Year: " + str(row[3]) + "\n"
                    "CGPA: " + str(row[4]) + "\n"
                    "Company: " + str(row[5]) + "\n"
                    "Status: " + str(row[6])
                )
                found = True
                break

        wb.close()

        if not found:
            result.delete("1.0", tk.END)
            result.insert(tk.END, "Student not found.")

    tk.Button(
        window,
        text="Search",
        width=15,
        command=search
    ).pack()


# ---------------- UPDATE STUDENT ----------------
def update_student():
    window = tk.Toplevel(dashboard)
    window.title("Update Student")
    window.geometry("450x400")

    tk.Label(
        window,
        text="Update Student",
        font=("Arial", 16, "bold")
    ).pack(pady=15)

    tk.Label(window, text="Student ID").pack()

    id_entry = tk.Entry(window)
    id_entry.pack(pady=5)

    tk.Label(window, text="New Company Name").pack()

    company_entry = tk.Entry(window)
    company_entry.pack(pady=5)

    tk.Label(window, text="New Placement Status").pack()

    status_entry = tk.Entry(window)
    status_entry.pack(pady=5)

    def update():
        student_id = id_entry.get().strip()
        company = company_entry.get().strip()
        status = status_entry.get().strip()

        wb = openpyxl.load_workbook(FILE)
        ws = wb.active

        found = False

        for row in ws.iter_rows(min_row=2):
            if str(row[0].value) == student_id:
                row[5].value = company
                row[6].value = status
                found = True
                break

        if found:
            wb.save(FILE)
            messagebox.showinfo(
                "Success",
                "Student updated successfully!"
            )
            window.destroy()
        else:
            messagebox.showerror(
                "Error",
                "Student not found."
            )

        wb.close()

    tk.Button(
        window,
        text="Update",
        width=15,
        command=update
    ).pack(pady=20)


# ---------------- DELETE STUDENT ----------------
def delete_student():
    window = tk.Toplevel(dashboard)
    window.title("Delete Student")
    window.geometry("400x250")

    tk.Label(
        window,
        text="Delete Student",
        font=("Arial", 16, "bold")
    ).pack(pady=20)

    tk.Label(window, text="Student ID").pack()

    entry = tk.Entry(window)
    entry.pack(pady=10)

    def delete():
        student_id = entry.get().strip()

        wb = openpyxl.load_workbook(FILE)
        ws = wb.active

        found = False

        for row in range(2, ws.max_row + 1):
            if str(ws.cell(row, 1).value) == student_id:
                ws.delete_rows(row, 1)
                found = True
                break

        if found:
            wb.save(FILE)
            messagebox.showinfo(
                "Success",
                "Student deleted successfully!"
            )
            window.destroy()
        else:
            messagebox.showerror(
                "Error",
                "Student not found."
            )

        wb.close()

    tk.Button(
        window,
        text="Delete",
        width=15,
        command=delete
    ).pack(pady=15)


# ---------------- DASHBOARD ----------------
def open_dashboard():
    global dashboard

    login_window.destroy()

    dashboard = tk.Tk()
    dashboard.title("College Placement Management System")
    dashboard.geometry("500x550")

    tk.Label(
        dashboard,
        text="College Placement Management System",
        font=("Arial", 16, "bold")
    ).pack(pady=25)

    tk.Button(
        dashboard,
        text="Add Student",
        width=25,
        command=add_student
    ).pack(pady=8)

    tk.Button(
        dashboard,
        text="View Records",
        width=25,
        command=view_records
    ).pack(pady=8)

    tk.Button(
        dashboard,
        text="Search Student",
        width=25,
        command=search_student
    ).pack(pady=8)

    tk.Button(
        dashboard,
        text="Update Student",
        width=25,
        command=update_student
    ).pack(pady=8)

    tk.Button(
        dashboard,
        text="Delete Student",
        width=25,
        command=delete_student
    ).pack(pady=8)

    tk.Button(
        dashboard,
        text="Logout",
        width=25,
        command=dashboard.destroy
    ).pack(pady=8)

    dashboard.mainloop()


# ---------------- LOGIN ----------------
login_window = tk.Tk()
login_window.title("College Placement Management System")
login_window.geometry("500x400")

tk.Label(
    login_window,
    text="College Placement Management System",
    font=("Arial", 16, "bold")
).pack(pady=30)

tk.Label(login_window, text="Username").pack()

username = tk.Entry(login_window)
username.pack(pady=5)

tk.Label(login_window, text="Password").pack()

password = tk.Entry(login_window, show="*")
password.pack(pady=5)


def login():
    if username.get() == "admin" and password.get() == "1234":
        messagebox.showinfo(
            "Login",
            "Login Successful!"
        )
        open_dashboard()
    else:
        messagebox.showerror(
            "Login",
            "Invalid Username or Password"
        )


tk.Button(
    login_window,
    text="Login",
    command=login
).pack(pady=20)

login_window.mainloop()